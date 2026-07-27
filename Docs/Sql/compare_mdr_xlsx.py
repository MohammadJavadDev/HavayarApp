import re
import zipfile
import openpyxl

ref = r"c:\Users\Padidar\Downloads\Hts_Edms_Mdr_Report_20260630100831_235847.xlsx"
gen = r"d:\Projects\Havayar\HavayarApp\mdr_export_test.xlsx"


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    return wb["Sheet1"]


def merges(path):
    with zipfile.ZipFile(path) as z:
        sv = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        return sorted(re.findall(r'mergeCell ref="([^"]+)"', sv))


def row_values(ws, row, max_col=467):
    return {c: ws.cell(row, c).value for c in range(1, max_col + 1) if ws.cell(row, c).value not in (None, "")}


def header_row(ws, row):
    return [ws.cell(row, c).value for c in range(1, 468)]


ref_ws = load(ref)
gen_ws = load(gen)

print("ref sheet: Sheet1 gen sheet: Sheet1")
print("ref freeze:", ref_ws.freeze_panes, "gen freeze:", gen_ws.freeze_panes)
print("merges equal:", merges(ref) == merges(gen))

r2_ref = header_row(ref_ws, 2)
r2_gen = header_row(gen_ws, 2)
hdr_diff = [(i + 1, a, b) for i, (a, b) in enumerate(zip(r2_ref, r2_gen)) if a != b]
print("row2 header diffs:", len(hdr_diff))
if hdr_diff[:5]:
    print(" sample:", hdr_diff[:5])

r1_ref = header_row(ref_ws, 1)
r1_gen = header_row(gen_ws, 1)
r1_diff = [(i + 1, a, b) for i, (a, b) in enumerate(zip(r1_ref, r1_gen)) if a != b]
print("row1 section diffs:", len(r1_diff))

# compare sample data row 3 for overlapping keys from json
sample_cols = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 17, 20, 21, 22, 24, 25, 26, 27, 28, 29, 30, 34, 37, 39, 41, 47, 48, 49, 50, 51, 52, 53, 54, 55, 59, 60, 61]
print("\nRow3 column compare:")
for c in sample_cols:
    rv = ref_ws.cell(3, c).value
    gv = gen_ws.cell(3, c).value
    if str(rv) != str(gv):
        print(f"  col {c}: ref={rv!r} gen={gv!r}")

# style compare A1
for label, ws in [("ref", ref_ws), ("gen", gen_ws)]:
    c = ws.cell(1, 1)
    print(f"{label} A1 fill={c.fill.start_color.rgb if c.fill else None} font={c.font.name}/{c.font.size}")
