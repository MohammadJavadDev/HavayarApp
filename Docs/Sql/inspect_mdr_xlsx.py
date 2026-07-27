import re
import zipfile

path = r"c:\Users\Padidar\Downloads\Hts_Edms_Mdr_Report_20260630100831_235847.xlsx"
with zipfile.ZipFile(path) as z:
    sv = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    pane = re.search(r"<pane[^>]*/>", sv)
    print("pane:", pane.group(0) if pane else "none")
    merges = re.findall(r'<mergeCell ref="([^"]+)"', sv)
    print("merge_count:", len(merges))
    print("merges:", merges[:5], "...", merges[-3:])
    wb = z.read("xl/workbook.xml").decode("utf-8")
    print("sheets:", re.findall(r'sheet name="([^"]+)"', wb))

try:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    print("active:", ws.title, "rows:", ws.max_row, "cols:", ws.max_column)
    print("freeze:", ws.freeze_panes)
    c = ws.cell(1, 1)
    print("A1 fill:", c.fill.start_color.rgb if c.fill else None, "font:", c.font.name, c.font.size, c.font.color.rgb if c.font and c.font.color else None)
    c2 = ws.cell(2, 1)
    print("A2 fill:", c2.fill.start_color.rgb if c2.fill else None)
    c34 = ws.cell(1, 34)
    print("AH1:", c34.value, "fill:", c34.fill.start_color.rgb if c34.fill else None)
    # borders on header
    b = c.border
    print("A1 border:", b.left.style, b.top.style, b.right.style, b.bottom.style)
    d3 = ws.cell(3, 1)
    print("A3 border:", d3.border.left.style if d3.border else None)
    print("A3 fill:", d3.fill.start_color.rgb if d3.fill and d3.fill.start_color else None)
    print("row1 height:", ws.row_dimensions[1].height)
    print("row2 height:", ws.row_dimensions[2].height)
    print("row3 height:", ws.row_dimensions[3].height)
    print("col A width:", ws.column_dimensions["A"].width)
    print("col B width:", ws.column_dimensions["B"].width)
except ImportError:
    print("openpyxl not installed")
